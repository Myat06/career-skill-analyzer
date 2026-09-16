"""Regression tests for docx extraction.

Table-based layouts are one of the most common resume template styles, and
reading document.paragraphs alone skipped every one of them -- a table-laid-out
resume extracted to nothing but the candidate's name. Nothing errored; the
rubric grader simply had no text to grade, so the report printed "no
improvements suggested" over a resume it had never actually read.
"""

import io

import pytest
from docx import Document

from app.resume.parser import UnsupportedResumeFormat, extract_text, extract_text_from_docx


def _docx_bytes(build) -> bytes:
    doc = Document()
    build(doc)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_plain_paragraphs_still_extract():
    data = _docx_bytes(lambda d: [d.add_paragraph("Jane Doe"), d.add_paragraph("Skills: Python")])
    text = extract_text_from_docx(data)
    assert "Jane Doe" in text and "Skills: Python" in text


def test_table_cell_text_is_extracted():
    def build(doc):
        doc.add_paragraph("Jane Doe")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Email"
        table.cell(0, 1).text = "jane@example.com"
        table.cell(1, 0).text = "Experience"
        table.cell(1, 1).text = "Data Intern - built dashboards."

    text = extract_text_from_docx(_docx_bytes(build))
    assert "jane@example.com" in text
    assert "Data Intern - built dashboards." in text


def test_document_order_is_preserved_across_paragraphs_and_tables():
    def build(doc):
        doc.add_paragraph("FIRST")
        doc.add_table(rows=1, cols=1).cell(0, 0).text = "SECOND"
        doc.add_paragraph("THIRD")

    text = extract_text_from_docx(_docx_bytes(build))
    assert text.index("FIRST") < text.index("SECOND") < text.index("THIRD")


def test_nested_table_text_is_extracted():
    def build(doc):
        outer = doc.add_table(rows=1, cols=1)
        inner = outer.cell(0, 0).add_table(rows=1, cols=1)
        inner.cell(0, 0).text = "nested-value"

    assert "nested-value" in extract_text_from_docx(_docx_bytes(build))


def test_header_text_is_extracted():
    # The rubric's first category grades contact details, which are routinely
    # placed in the header -- a resume shouldn't be marked down for "missing" a
    # phone number that was in the file all along.
    def build(doc):
        doc.add_paragraph("Jane Doe")
        doc.sections[0].header.paragraphs[0].text = "+62 812 0000 0000"

    assert "+62 812 0000 0000" in extract_text_from_docx(_docx_bytes(build))


def test_merged_cells_are_not_duplicated():
    def build(doc):
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "SPANNED"
        table.cell(0, 0).merge(table.cell(0, 1))

    assert extract_text_from_docx(_docx_bytes(build)).count("SPANNED") == 1


def test_empty_cells_do_not_produce_blank_lines():
    def build(doc):
        doc.add_paragraph("Jane Doe")
        doc.add_table(rows=2, cols=2)  # entirely empty table

    text = extract_text_from_docx(_docx_bytes(build))
    assert text.strip() == "Jane Doe"


def test_unsupported_extension_raises():
    with pytest.raises(UnsupportedResumeFormat):
        extract_text("resume.txt", b"whatever")
